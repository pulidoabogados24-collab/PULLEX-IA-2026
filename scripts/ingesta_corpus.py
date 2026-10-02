"""Carga el corpus propio de PULLEX (códigos, leyes, sentencias, plantillas, doctrina) en el índice
SQLite FTS5 que usa el chat (fuentes.py). No necesita claves de IA.

Uso (desde la carpeta del proyecto):

    # 1) Desde una carpeta de tu computador (PDF, DOCX, DOC, RTF, TXT, MD), con subcarpetas:
    python scripts/ingesta_corpus.py --carpeta "C:/Users/tu/Descargas/LEXCOL_CORPUS"

    # 2) Desde Google Drive con una cuenta de servicio (dependencias opcionales):
    #    pip install google-api-python-client google-auth
    #    GOOGLE_SERVICE_ACCOUNT_JSON=ruta/al/archivo.json  DRIVE_FOLDER_ID=<id de LEXCOL_CORPUS>
    python scripts/ingesta_corpus.py --drive

    # Ver qué se indexaría y qué se excluiría, sin escribir nada:
    python scripts/ingesta_corpus.py --carpeta ... --simular

    # Después de confirmar un documento en la fuente oficial:
    python scripts/ingesta_corpus.py --listar
    python scripts/ingesta_corpus.py --verificar 12          (queda VIGENTE_VERIFICADA hoy)
    python scripts/ingesta_corpus.py --estado 7 DEROGADA

Reglas (detalle en docs/11-MOTOR-DE-FUENTES.md):
- Todo documento entra como PENDIENTE_VERIFICAR; solo una persona lo marca verificado.
- NO se indexa nada que parezca de clientes o personas reales: carpetas o archivos con
  "CLIENTE" o "EXPEDIENTE", carpetas cuyo nombre es solo un nombre propio en mayúsculas,
  archivos "NOMBRE APELLIDO + tipo de escrito" (p. ej. "JUAN PEREZ TUTELA.pdf"), y lo que
  esté en la lista de exclusión (PULLEX_CORPUS_EXCLUIR o corpus/excluir.txt).
"""
import argparse
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
import fuentes  # noqa: E402

EXTENSIONES = {".pdf", ".docx", ".doc", ".rtf", ".txt", ".md"}
MIME_DRIVE = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/msword": ".doc",          # Word 97-2003: necesita un conversor (ver extraer_doc)
    "application/rtf": ".rtf",
    "text/rtf": ".rtf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.google-apps.document": ".gdoc",   # se exporta como texto plano
}
CARPETA_DRIVE = "application/vnd.google-apps.folder"


# ----------------------------------------------------------- extracción --
def _pdf_pypdf(datos: bytes):
    import pypdf  # opcional
    lector = pypdf.PdfReader(io.BytesIO(datos))
    return [(f"pág. {i}", p.extract_text() or "") for i, p in enumerate(lector.pages, 1)]


def _pdf_pdftotext(datos: bytes):
    if not shutil.which("pdftotext"):
        raise RuntimeError("no hay pypdf ni pdftotext: instala uno de los dos (pip install pypdf)")
    r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", "-", "-"], input=datos,
                       capture_output=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError("pdftotext falló: " + r.stderr.decode("utf-8", "ignore")[:200])
    paginas = r.stdout.decode("utf-8", "ignore").split("\f")
    return [(f"pág. {i}", t) for i, t in enumerate(paginas, 1)]


def extraer_pdf(datos: bytes):
    try:
        return _pdf_pypdf(datos)
    except ImportError:
        return _pdf_pdftotext(datos)


def _docx_xml(datos: bytes):
    """Texto de un .docx leyendo directamente word/document.xml (sin python-docx)."""
    import re
    import zipfile
    with zipfile.ZipFile(io.BytesIO(datos)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    xml = re.sub(r"</w:p>", "\n", xml)
    return [("", re.sub(r"<[^>]+>", "", xml))]


def extraer_docx(datos: bytes):
    try:
        import docx  # python-docx, opcional
    except ImportError:
        return _docx_xml(datos)
    try:
        d = docx.Document(io.BytesIO(datos))
    except Exception:
        # .docx mínimo o generado por otra herramienta que python-docx no abre: se lee el XML.
        return _docx_xml(datos)
    partes = [p.text for p in d.paragraphs]
    for tabla in d.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(c.text for c in fila.cells))
    return [("", "\n".join(partes))]


_RTF_DESTINOS = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "generator", "themedata",
                 "datastore", "latentstyles", "listtable", "listoverridetable", "rsidtbl", "xmlnstbl", "object", "fldinst"}
_RTF_SALTOS = {"par": "\n", "line": "\n", "sect": "\n\n", "page": "\n\n", "tab": "\t", "cell": " | ", "row": "\n",
               "emdash": "—", "endash": "–", "lquote": "‘", "rquote": "’", "ldblquote": "“", "rdblquote": "”", "bullet": "•"}


def rtf_a_texto(datos: bytes) -> str:
    """Texto de un .rtf sin dependencias: quita las palabras de control y las tablas de fuentes o
    estilos, y traduce \\'hh (página de códigos 1252) y \\uN. No interpreta campos ni imágenes."""
    import re
    rtf = datos.decode("latin-1", "ignore")
    token = re.compile(r"\\([a-z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-fA-F]{2})|\\([^a-z])|([{}])|[\r\n]+|(.)", re.I | re.S)
    pila, ignorar, saltar, salida = [], False, 0, []
    for m in token.finditer(rtf):
        palabra, arg, hexa, simbolo, llave, car = m.groups()
        if llave == "{":
            pila.append(ignorar)
        elif llave == "}":
            ignorar = pila.pop() if pila else False
        elif simbolo is not None:
            if simbolo == "*":
                ignorar = True                       # destino opcional: se ignora entero
            elif not ignorar and simbolo in "\\{}":
                salida.append(simbolo)
            elif not ignorar and simbolo == "~":
                salida.append(" ")
        elif palabra is not None:
            if palabra in _RTF_DESTINOS:
                ignorar = True
            elif ignorar:
                continue
            elif palabra in _RTF_SALTOS:
                salida.append(_RTF_SALTOS[palabra])
            elif palabra == "u" and arg:
                n = int(arg)
                salida.append(chr(n + 65536 if n < 0 else n))
                saltar = 1                           # tras \\uN viene un carácter de reemplazo que no se usa
        elif hexa is not None:
            if saltar:
                saltar -= 1
            elif not ignorar:
                salida.append(bytes([int(hexa, 16)]).decode("cp1252", "ignore"))
        elif car is not None and not ignorar:
            if saltar:
                saltar -= 1
            else:
                salida.append(car)
    return re.sub(r"[ \t]+\n", "\n", "".join(salida)).strip()


def extraer_doc(datos: bytes):
    """Texto de un .doc (Word 97-2003). No hay lector puro en Python: usa el conversor que esté
    instalado (antiword, catdoc o LibreOffice). Sin ninguno, falla con un mensaje que lo dice."""
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        entrada = Path(tmp) / "entrada.doc"
        entrada.write_bytes(datos)
        for orden in (["antiword", "-m", "UTF-8.txt", "-w", "0", str(entrada)], ["catdoc", "-d", "utf-8", "-w", str(entrada)]):
            if shutil.which(orden[0]):
                r = subprocess.run(orden, capture_output=True, timeout=120)
                if r.returncode == 0 and r.stdout.strip():
                    return [("", r.stdout.decode("utf-8", "ignore"))]
        oficina = shutil.which("soffice") or shutil.which("libreoffice")
        if oficina:
            r = subprocess.run([oficina, "--headless", "--convert-to", "txt:Text (encoded):UTF8", "--outdir", tmp, str(entrada)],
                               capture_output=True, timeout=300)
            salida = Path(tmp) / "entrada.txt"
            if salida.is_file():
                return [("", salida.read_text(encoding="utf-8-sig", errors="ignore"))]
            raise RuntimeError("LibreOffice no pudo convertir el .doc: " + r.stderr.decode("utf-8", "ignore")[:160])
    raise RuntimeError("para leer archivos .doc instala un conversor: antiword, catdoc o LibreOffice (soffice)")


def extraer(nombre: str, datos: bytes):
    ext = Path(nombre).suffix.lower()
    if ext == ".pdf":
        return extraer_pdf(datos)
    if ext == ".docx":
        return extraer_docx(datos)
    if ext == ".doc":
        return extraer_doc(datos)
    if ext == ".rtf":
        return [("", rtf_a_texto(datos))]
    if ext in (".txt", ".md", ".gdoc"):
        return [("", datos.decode("utf-8", "ignore"))]
    raise ValueError(f"extensión no admitida: {ext}")


# ------------------------------------------------------------- fuentes de archivos --
def recorrer_carpeta(base: Path):
    """(ruta_relativa, carpeta, leer(), fecha_iso, origen, url) para cada archivo admitido."""
    for p in sorted(base.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in EXTENSIONES or p.name.startswith((".", "~$")):
            continue
        rel = p.relative_to(base).as_posix()
        fecha = datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat(timespec="seconds")
        yield rel, p.parent.relative_to(base).as_posix(), p.read_bytes, fecha, "local:" + rel, None


def _servicio_drive():
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
    except ImportError:
        sys.exit("Para leer Google Drive instala las dependencias opcionales:\n"
                 "    pip install google-api-python-client google-auth")
    cred = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "")
    if not cred or not os.path.isfile(cred):
        sys.exit("Define GOOGLE_SERVICE_ACCOUNT_JSON con la ruta al archivo .json de la cuenta de servicio.")
    c = service_account.Credentials.from_service_account_file(
        cred, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=c, cache_discovery=False)


def recorrer_drive(servicio, carpeta_id: str, ruta: str = "", patrones=None, reporte=None):
    """Recorre la carpeta y sus subcarpetas. Las carpetas excluidas ni siquiera se abren."""
    token = None
    while True:
        r = servicio.files().list(
            q=f"'{carpeta_id}' in parents and trashed=false",
            fields="nextPageToken, files(id, name, mimeType, modifiedTime)",
            pageSize=200, pageToken=token, supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        for f in r.get("files", []):
            rel = f"{ruta}/{f['name']}" if ruta else f["name"]
            if f["mimeType"] == CARPETA_DRIVE:
                motivo = fuentes.es_excluido(rel + "/x.txt", patrones)
                if motivo and reporte is not None:
                    reporte["excluidos"].append({"ruta": rel + "/", "motivo": motivo})
                    continue
                yield from recorrer_drive(servicio, f["id"], rel, patrones, reporte)
                continue
            ext = MIME_DRIVE.get(f["mimeType"])
            if not ext:
                # nada se salta en silencio: lo que no se puede leer queda en el reporte con su motivo
                if reporte is not None:
                    motivo = ("acceso directo: no se sigue; comparte la carpeta de destino" if f["mimeType"].endswith(".shortcut")
                              else "formato no admitido (" + f["mimeType"].split("/")[-1].split(".")[-1] + ")")
                    reporte.setdefault("omitidos", []).append({"ruta": rel, "motivo": motivo})
                continue
            nombre = f["name"] if ext == ".gdoc" or f["name"].lower().endswith(ext) else f["name"] + ext

            def leer(fid=f["id"], mime=f["mimeType"]):
                if mime == "application/vnd.google-apps.document":
                    return servicio.files().export(fileId=fid, mimeType="text/plain").execute()
                return servicio.files().get_media(fileId=fid, supportsAllDrives=True).execute()

            yield (f"{ruta}/{nombre}" if ruta else nombre, ruta, leer, f.get("modifiedTime"),
                   "drive:" + f["id"], f"https://drive.google.com/file/d/{f['id']}/view")
        token = r.get("nextPageToken")
        if not token:
            break


# ------------------------------------------------------------------- ingesta --
def ingerir(origenes, db_ruta: str, simular: bool = False, patrones=None, reporte=None) -> dict:
    reporte = reporte if reporte is not None else {"excluidos": []}
    reporte.update(indexados=[], sin_cambios=0, errores=[])
    con = None if simular else fuentes.abrir(db_ruta)
    try:
        for rel, carpeta, leer, fecha, origen, url in origenes:
            motivo = fuentes.es_excluido(rel, patrones)
            if motivo:
                reporte["excluidos"].append({"ruta": rel, "motivo": motivo})
                continue
            if con is not None and origen.startswith("drive:") and con.execute(
                    "SELECT 1 FROM fuentes WHERE origen=?", (origen[6:],)).fetchone():
                # ya lo indexó la biblioteca (scripts/indexar_biblioteca.py) con el id de Drive como origen
                reporte["ya_en_biblioteca"] = reporte.get("ya_en_biblioteca", 0) + 1
                continue
            if simular:
                c = fuentes.clasificar_nombre(rel, carpeta)
                reporte["indexados"].append({"ruta": rel, "tipo": c["tipo"], "estado": fuentes.estado_inicial(fecha)})
                continue
            try:
                datos = leer()
                paginas = extraer(rel, datos)
                r = fuentes.indexar(con, origen=origen, nombre=Path(rel).name, paginas=paginas, carpeta=carpeta,
                                    fecha_archivo=fecha, url=url, sha256=hashlib.sha256(datos).hexdigest())
            except Exception as e:  # un archivo dañado no detiene la carga
                reporte["errores"].append({"ruta": rel, "error": f"{type(e).__name__}: {str(e)[:160]}"})
                continue
            if r["accion"] == "sin_cambios":
                reporte["sin_cambios"] += 1
            else:
                reporte["indexados"].append({"ruta": rel, "tipo": r["tipo"], "estado": r["estado"],
                                             "fragmentos": r["fragmentos"], "accion": r["accion"]})
                if r["fragmentos"] == 0:
                    reporte["errores"].append({"ruta": rel, "error": "sin texto extraíble (¿PDF escaneado? necesita OCR)"})
    finally:
        if con is not None:
            con.close()
    return reporte


def imprimir_reporte(rep: dict):
    print(f"\nIndexados: {len(rep['indexados'])}   Sin cambios: {rep.get('sin_cambios', 0)}   "
          f"Excluidos: {len(rep['excluidos'])}   Con error: {len(rep['errores'])}")
    for x in rep["indexados"]:
        print(f"  + [{x['tipo']:<9}] {x['estado']:<20} {x['ruta']}" + (f"  ({x['fragmentos']} fragmentos)" if "fragmentos" in x else ""))
    for x in rep["excluidos"]:
        print(f"  - EXCLUIDO  {x['ruta']}  → {x['motivo']}")
    for x in rep["errores"]:
        print(f"  ! ERROR     {x['ruta']}  → {x['error']}")
    if rep.get("ya_en_biblioteca"):
        print(f"  = {rep['ya_en_biblioteca']} ya estaban indexados por la biblioteca (no se duplican)")
    omitidos = rep.get("omitidos", [])
    if omitidos:
        from collections import Counter
        print(f"\nOmitidos (no se pueden leer): {len(omitidos)}")
        for motivo, n in Counter(x["motivo"] for x in omitidos).most_common():
            print(f"  · {n:>5}  {motivo}")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--carpeta", help="carpeta local con el corpus")
    ap.add_argument("--drive", action="store_true", help="leer desde Google Drive (DRIVE_FOLDER_ID)")
    ap.add_argument("--db", default=None, help="ruta del índice (por defecto PULLEX_CORPUS_DB o corpus/corpus.db)")
    ap.add_argument("--simular", action="store_true", help="solo mostrar qué se indexaría y qué se excluiría")
    ap.add_argument("--excluir", action="append", default=[], help="patrón adicional de exclusión (repetible)")
    ap.add_argument("--listar", action="store_true", help="listar las fuentes indexadas")
    ap.add_argument("--verificar", type=int, metavar="ID", help="marcar VIGENTE_VERIFICADA con fecha de hoy")
    ap.add_argument("--estado", nargs=2, metavar=("ID", "ESTADO"), help="fijar estado: " + ", ".join(fuentes.ESTADOS))
    ap.add_argument("--json", action="store_true", help="imprimir el reporte en JSON")
    a = ap.parse_args(argv)
    db_ruta = a.db or fuentes.ruta_db()
    patrones = fuentes.patrones_exclusion(a.excluir)

    if a.listar or a.verificar or a.estado:
        con = fuentes.abrir(db_ruta)
        if a.verificar:
            print("OK" if fuentes.marcar_estado(con, a.verificar, "VIGENTE_VERIFICADA") else "No existe ese ID")
        if a.estado:
            print("OK" if fuentes.marcar_estado(con, int(a.estado[0]), a.estado[1].upper()) else "No existe ese ID")
        if a.listar:
            for f in con.execute("SELECT s.*, (SELECT COUNT(*) FROM fragmentos WHERE fuente_id=s.id) n "
                                 "FROM fuentes s ORDER BY s.tipo, s.titulo"):
                ef = fuentes.estado_efectivo(f["estado_vigencia"], f["verificado_en"])
                print(f"{f['id']:>4}  [{f['tipo']:<9}] {ef:<20} {f['titulo']}  ({f['n']} frag., archivo {str(f['fecha_archivo'] or '')[:10]})")
        con.close()
        return 0

    reporte = {"excluidos": []}
    if a.carpeta:
        base = Path(a.carpeta).expanduser()
        if not base.is_dir():
            sys.exit(f"No existe la carpeta: {base}")
        origenes = recorrer_carpeta(base)
    elif a.drive:
        carpeta_id = os.getenv("DRIVE_FOLDER_ID", "")
        if not carpeta_id:
            sys.exit("Define DRIVE_FOLDER_ID con el id de la carpeta LEXCOL_CORPUS en Google Drive.")
        origenes = recorrer_drive(_servicio_drive(), carpeta_id, "", patrones, reporte)
    else:
        ap.print_help()
        return 2
    rep = ingerir(origenes, db_ruta, simular=a.simular, patrones=patrones, reporte=reporte)
    if a.json:
        print(json.dumps(rep, ensure_ascii=False, indent=1))
    else:
        imprimir_reporte(rep)
        if not a.simular:
            print(f"\nÍndice: {db_ruta}. Todo lo nuevo quedó PENDIENTE_VERIFICAR: confirma en la fuente oficial "
                  "y usa --verificar ID.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
